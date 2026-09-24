import streamlit as st
import pandas as pd
import json
import os
import io
from datetime import datetime
from dotenv import load_dotenv, set_key
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Muat variabel lingkungan dari .env
ENV_PATH = os.path.join(os.path.dirname(__file__), ".env")
load_dotenv(ENV_PATH)

# Kredensial Administrator
ADMIN_USER = "alex"
ADMIN_PASS = "Wafie"

# Helper untuk konfigurasi & Secrets (Kompatibel Local + Streamlit Cloud)
def get_config(key, default=""):
    try:
        if hasattr(st, "secrets") and key in st.secrets:
            return str(st.secrets[key])
    except Exception:
        pass
    return os.getenv(key, default)

# Konfigurasi Baku
DEFAULT_API_KEY = get_config("GEMINI_API_KEY", "")
LOG_DIR = os.path.join(os.path.dirname(__file__), "data")
LOG_FILE = os.path.join(LOG_DIR, "audit_history.json")

# Inisialisasi Google GenAI SDK
try:
    from google import genai
    HAS_NEW_GENAI = True
except ImportError:
    HAS_NEW_GENAI = False
    import google.generativeai as legacy_genai

# ==============================================================================
# FUNGSI PERSISTENSI LOG AUDIT
# ==============================================================================
def load_audit_logs():
    if not os.path.exists(LOG_DIR):
        os.makedirs(LOG_DIR, exist_ok=True)
    if not os.path.exists(LOG_FILE):
        return []
    try:
        with open(LOG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []

def save_audit_log(entry):
    logs = load_audit_logs()
    logs.insert(0, entry)
    with open(LOG_FILE, "w", encoding="utf-8") as f:
        json.dump(logs, f, ensure_ascii=False, indent=2)

def delete_audit_log(log_id):
    logs = load_audit_logs()
    logs = [l for l in logs if l.get("id") != log_id]
    with open(LOG_FILE, "w", encoding="utf-8") as f:
        json.dump(logs, f, ensure_ascii=False, indent=2)

def clear_all_audit_logs():
    with open(LOG_FILE, "w", encoding="utf-8") as f:
        json.dump([], f)

# ==============================================================================
# KONFIGURASI DAN DESAIN SISTEM ENTERPRISE
# ==============================================================================
st.set_page_config(
    page_title="Sistem Surveilans Epidemiologi Terpadu",
    page_icon="🏢",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap" rel="stylesheet">

<style>
    /* Font Global */
    html, body, [class*="css"], .stMarkdown {
        font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, sans-serif !important;
    }
    
    .block-container {
        padding-top: 3.5rem !important;
        padding-bottom: 4rem !important;
        max-width: 1140px !important;
    }
    
    /* Tombol Utama Universal - Modern Emerald */
    div.stButton > button[kind="primary"],
    div.stButton > button[type="primary"] {
        background-color: #0f766e !important;
        color: #ffffff !important;
        border: 1px solid #0d9488 !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
        padding: 0.5rem 1.25rem !important;
        transition: all 0.2s ease !important;
    }
    div.stButton > button[kind="primary"]:hover,
    div.stButton > button[type="primary"]:hover {
        background-color: #115e59 !important;
        border-color: #14b8a6 !important;
        box-shadow: 0 4px 12px rgba(15, 118, 110, 0.25) !important;
    }
    div.stButton > button[kind="secondary"] {
        border-radius: 6px !important;
        font-weight: 500 !important;
    }
    
    /* Top Bar Navigasi */
    .top-navbar {
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-bottom: 1px solid rgba(255, 255, 255, 0.12);
        padding-bottom: 16px;
        margin-bottom: 28px;
    }
    .brand-title {
        font-size: 20px;
        font-weight: 700;
        letter-spacing: -0.3px;
        margin: 0;
        color: #ffffff !important;
        line-height: 1.3;
    }
    .brand-subtitle {
        font-size: 13px;
        color: #94a3b8 !important;
        margin-top: 4px;
        line-height: 1.4;
    }
    .status-pill {
        font-size: 11px;
        font-weight: 600;
        letter-spacing: 0.5px;
        padding: 5px 12px;
        border-radius: 4px;
        background: rgba(16, 185, 129, 0.15);
        color: #10b981 !important;
        border: 1px solid rgba(16, 185, 129, 0.3);
        white-space: nowrap;
    }
    
    /* Hero Banner */
    .hero-card {
        background: linear-gradient(180deg, rgba(30, 41, 59, 0.7) 0%, rgba(15, 23, 42, 0.85) 100%);
        border: 1px solid rgba(255, 255, 255, 0.1);
        border-radius: 10px;
        padding: 26px 30px;
        margin-bottom: 26px;
    }
    .hero-heading {
        font-size: 20px;
        font-weight: 700;
        color: #ffffff !important;
        margin-bottom: 8px;
    }
    .hero-body {
        font-size: 14px;
        color: #cbd5e1 !important;
        line-height: 1.65;
    }
    
    /* Step Workflow Grid */
    .workflow-card {
        background: rgba(30, 41, 59, 0.5);
        border: 1px solid rgba(255, 255, 255, 0.08);
        border-radius: 8px;
        padding: 20px;
        height: 100%;
    }
    .workflow-badge {
        font-size: 11px;
        font-weight: 700;
        letter-spacing: 0.5px;
        color: #14b8a6 !important;
        text-transform: uppercase;
        margin-bottom: 8px;
    }
    .workflow-title {
        font-size: 15px;
        font-weight: 600;
        color: #f8fafc !important;
        margin-bottom: 6px;
    }
    .workflow-desc {
        font-size: 13px;
        color: #94a3b8 !important;
        line-height: 1.55;
    }
    
    /* Info Card */
    .info-container {
        background: rgba(30, 41, 59, 0.4);
        border: 1px solid rgba(255, 255, 255, 0.08);
        border-radius: 8px;
        padding: 22px 24px;
        margin-top: 14px;
    }
    
    /* Sidebar */
    section[data-testid="stSidebar"] {
        border-right: 1px solid rgba(255, 255, 255, 0.08);
    }
    .sidebar-brand {
        padding-bottom: 12px;
        border-bottom: 1px solid rgba(255, 255, 255, 0.1);
        margin-bottom: 16px;
    }
    
    /* Dokumen Output Fisik */
    .official-paper {
        background-color: #ffffff !important;
        color: #0f172a !important;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        padding: 40px 48px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.06);
        font-size: 14px;
        line-height: 1.7;
    }
    .official-paper h1, .official-paper h2, .official-paper h3, .official-paper h4 {
        color: #0f172a !important;
    }
</style>
""", unsafe_allow_html=True)

# Inisialisasi State Aplikasi
if "df_data" not in st.session_state:
    st.session_state.df_data = None
if "laporan_teks" not in st.session_state:
    st.session_state.laporan_teks = ""
if "is_admin" not in st.session_state:
    st.session_state.is_admin = False
if "api_key" not in st.session_state:
    st.session_state.api_key = DEFAULT_API_KEY
if "instansi" not in st.session_state:
    st.session_state.instansi = get_config("DEFAULT_INSTANSI", "Dinas Kesehatan Kabupaten/Kota")
if "bidang" not in st.session_state:
    st.session_state.bidang = get_config("DEFAULT_BIDANG", "Bidang Pencegahan dan Pengendalian Penyakit (P2P)")
if "model_name" not in st.session_state:
    st.session_state.model_name = get_config("DEFAULT_MODEL", "gemini-3.5-flash")
if "active_page" not in st.session_state:
    st.session_state.active_page = "tutorial"

# ==============================================================================
# SIDEBAR RESMI DINAS
# ==============================================================================
with st.sidebar:
    st.markdown("""
    <div class="sidebar-brand">
        <div style="font-size: 15px; font-weight: 700; color: #ffffff;">Dinas Kesehatan</div>
        <div style="font-size: 12px; color: #94a3b8; margin-top: 2px;">Unit Surveilans Epidemiologi</div>
    </div>
    """, unsafe_allow_html=True)
    
    st.caption("MENU UTAMA")
    
    col_nav1, col_nav2 = st.columns(2)
    with col_nav1:
        if st.button("Panduan", use_container_width=True, type="primary" if st.session_state.active_page == "tutorial" else "secondary"):
            st.session_state.active_page = "tutorial"
            st.rerun()
    with col_nav2:
        if st.button("Workspace", use_container_width=True, type="primary" if st.session_state.active_page == "workspace" else "secondary"):
            st.session_state.active_page = "workspace"
            st.rerun()
            
    if st.session_state.is_admin:
        if st.button("Portal Admin (Alex)", use_container_width=True, type="primary" if st.session_state.active_page == "admin" else "secondary"):
            st.session_state.active_page = "admin"
            st.rerun()
            
    st.divider()
    
    # Autentikasi Admin
    if not st.session_state.is_admin:
        st.caption("AKSES KHUSUS")
        with st.expander("Login Administrator", expanded=False):
            with st.form("admin_login_box"):
                u_in = st.text_input("Username", placeholder="alex")
                p_in = st.text_input("Password", type="password", placeholder="••••••")
                if st.form_submit_button("Masuk", use_container_width=True, type="primary"):
                    if u_in == ADMIN_USER and p_in == ADMIN_PASS:
                        st.session_state.is_admin = True
                        st.session_state.active_page = "admin"
                        st.rerun()
                    else:
                        st.error("Kredensial tidak valid.")
    else:
        st.caption("STATUS OTENTIKASI")
        st.success(f"Masuk: **{ADMIN_USER}** (Admin)")
        if st.button("Keluar (Logout)", use_container_width=True):
            st.session_state.is_admin = False
            st.session_state.active_page = "tutorial"
            st.rerun()
            
    st.write("")
    st.caption("SISE v2.4 Enterprise Edition")

# ==============================================================================
# TOP BAR RESMI
# ==============================================================================
status_pill_text = "Sistem Operasional Aktif"
if st.session_state.is_admin:
    status_pill_text = "Administrator: Alex"

st.markdown(f"""
<div class="top-navbar">
    <div>
        <div class="brand-title">Sistem Laporan Naratif Surveilans Kesehatan</div>
        <div class="brand-subtitle">{st.session_state.instansi} &bull; {st.session_state.bidang}</div>
    </div>
    <div>
        <span class="status-pill">{status_pill_text}</span>
    </div>
</div>
""", unsafe_allow_html=True)

# ==============================================================================
# HALAMAN 1: PANDUAN & PENGENALAN SISTEM (INFORMATIF & MUDAH DIBACA)
# ==============================================================================
if st.session_state.active_page == "tutorial":
    
    # Hero Card
    st.markdown("""
    <div class="hero-card">
        <div class="hero-heading">Sistem Informasi Surveilans Epidemiologi Terpadu</div>
        <div class="hero-body">
            Platform resmi Dinas Kesehatan untuk mengonversi rekapan data angka mingguan Puskesmas 
            (kasus DBD, Hipertensi, Stunting, ISPA) menjadi <strong>Draf Laporan Naratif Resmi</strong> 
            lengkap dengan analisis hotspot, faktor risiko lingkungan, dan rekomendasi tindakan intervensi lapangan.
        </div>
    </div>
    """, unsafe_allow_html=True)
    
    # Tombol Masuk Workspace
    col_cta1, col_cta2 = st.columns([2, 1])
    with col_cta1:
        st.write("##### Siap menyusun laporan surveilans periode ini?")
    with col_cta2:
        if st.button("Buka Halaman Analisis Data ➔", type="primary", use_container_width=True):
            st.session_state.active_page = "workspace"
            st.rerun()
            
    st.divider()
    
    # 1. Alur 3 Langkah Kerja Cepat
    st.markdown("#### Alur Penggunaan Sederhana (3 Langkah)")
    st.caption("Sistem dirancang praktis agar staf dapat langsung menyelesaikan laporan dalam hitungan detik:")
    
    col_w1, col_w2, col_w3 = st.columns(3)
    
    with col_w1:
        st.markdown("""
        <div class="workflow-card">
            <div class="workflow-badge">Langkah 01</div>
            <div class="workflow-title">Unggah Rekapan Kasus</div>
            <div class="workflow-desc">
                Unggah file tabel (CSV atau Excel) berisi data kasus mingguan per Puskesmas. Anda juga bisa menguji langsung dengan tombol <em>"Gunakan Data Contoh"</em>.
            </div>
        </div>
        """, unsafe_allow_html=True)
        
    with col_w2:
        st.markdown("""
        <div class="workflow-card">
            <div class="workflow-badge">Langkah 02</div>
            <div class="workflow-title">Pilih Format & Generate</div>
            <div class="workflow-desc">
                Pilih format naskah yang diinginkan (Laporan Kedinasan Lengkap atau Ringkasan Eksekutif Pimpinan), lalu klik <strong>"Buat Laporan Otomatis"</strong>.
            </div>
        </div>
        """, unsafe_allow_html=True)
        
    with col_w3:
        st.markdown("""
        <div class="workflow-card">
            <div class="workflow-badge">Langkah 03</div>
            <div class="workflow-title">Review & Unduh Dokumen</div>
            <div class="workflow-desc">
                Baca hasil analisis pada lembar dokumen resmi, lakukan koreksi teks jika ada catatan khusus, lalu unduh dalam format <strong>Word (.docx)</strong> atau <strong>Teks (.txt)</strong>.
            </div>
        </div>
        """, unsafe_allow_html=True)

    st.write("")
    st.divider()
    
    # 2. Panduan Format Tabel yang Didukung
    st.markdown("#### Standar Format Tabel Data yang Didukung")
    st.caption("Sistem dapat membaca file CSV / Excel dengan struktur kolom fleksibel seperti contoh di bawah:")
    
    sample_preview_table = pd.DataFrame({
        "Puskesmas / Wilayah": ["Puskesmas Gambir", "Puskesmas Kebon Jeruk", "Puskesmas Cakung", "Puskesmas Penjaringan"],
        "Kasus DBD": [18, 32, 54, 41],
        "Kasus Hipertensi": [145, 230, 420, 350],
        "Stunting Baru": [4, 9, 16, 14],
        "Kasus ISPA": [210, 340, 560, 480],
        "Status Wilayah": ["Waspada", "Siaga", "KLB/Siaga", "Siaga"]
    })
    st.dataframe(sample_preview_table, use_container_width=True, hide_index=True)
    
    st.write("")
    
    # 3. Struktur Dokumen Resmi yang Dihasilkan
    st.markdown("#### Struktur Naskah Laporan yang Dihasilkan")
    st.caption("Draf laporan naratif otomatis memuat 5 bab standar Tata Naskah Dinas Kesehatan:")
    
    col_b1, col_b2 = st.columns(2)
    with col_b1:
        st.markdown("""
        - **Bab I. Latar Belakang & Ruang Lingkup**  
          Dasar hukum Sistem Kewaspadaan Dini dan Respon (SKDR) serta cakupan wilayah faskes pelapor.
        - **Bab II. Analisis Situasi & Pemetaan Hotspot**  
          Ringkasan agregat kasus dan pemetaan wilayah *Red Zone* (kasus tertinggi) vs wilayah aman.
        - **Bab III. Identifikasi Faktor Risiko Lapangan**  
          Analisis determinan sanitasi, cuaca, mobilitas penduduk, dan potensi eskalasi KLB.
        """)
    with col_b2:
        st.markdown("""
        - **Bab IV. Rekomendasi Tindakan & Intervensi**  
          Respon taktis cepat (PE 1x24 jam, fogging/larvasidasi, droping logistik) dan pencegahan jangka menengah.
        - **Bab V. Kesimpulan & Lembar Pengesahan**  
          Ringkasan eksekutif 1 paragraf serta slot tanggal & tanda tangan penanggung jawab resmi.
        """)
        
    st.write("")
    col_btm1, col_btm2, col_btm3 = st.columns([1, 2, 1])
    with col_btm2:
        if st.button("Mulai Olah Data Sekarang ➔", type="primary", use_container_width=True):
            st.session_state.active_page = "workspace"
            st.rerun()

# ==============================================================================
# HALAMAN 2: WORKSPACE GENERATOR LAPORAN
# ==============================================================================
elif st.session_state.active_page == "workspace":
    
    col_b1, col_b2 = st.columns([2, 1])
    with col_b1:
        if st.button("← Kembali ke Panduan & Pengenalan Sistem", use_container_width=False):
            st.session_state.active_page = "tutorial"
            st.rerun()
    with col_b2:
        st.write("")
        
    st.write("### Halaman Kerja Olah Data Surveilans")
    st.caption("Unggah file rekapan kasus untuk menghasilkan draf laporan naratif kedinasan resmi.")
    
    # 1. Upload File & Simulasi Data
    col_up, col_smp = st.columns([3, 1])
    with col_up:
        file_input = st.file_uploader(
            "Pilih file surveilans (CSV atau Excel):",
            type=["csv", "xlsx", "xls"],
            help="Unggah rekapitulasi mingguan puskesmas."
        )
    with col_smp:
        st.write("Data Simulasi:")
        if st.button("Gunakan Data Contoh", use_container_width=True):
            st.session_state.df_data = pd.DataFrame({
                "Puskesmas": ["Puskesmas Gambir", "Puskesmas Kebon Jeruk", "Puskesmas Cilandak", "Puskesmas Tebet", "Puskesmas Cakung", "Puskesmas Penjaringan", "Puskesmas Tanjung Priok", "Puskesmas Menteng"],
                "Kecamatan": ["Gambir", "Kebon Jeruk", "Cilandak", "Tebet", "Cakung", "Penjaringan", "Tanjung Priok", "Menteng"],
                "Kasus_DBD": [18, 32, 12, 27, 54, 41, 38, 9],
                "Kasus_Hipertensi": [145, 230, 180, 310, 420, 350, 290, 110],
                "Kasus_Stunting_Baru": [4, 9, 3, 8, 16, 14, 12, 2],
                "Kasus_ISPA": [210, 340, 195, 420, 560, 480, 390, 140],
                "Populasi_Beresiko": [12500, 24000, 18500, 29000, 41000, 36000, 33000, 9500],
                "Status_Kewaspadaan": ["Waspada", "Siaga", "Aman", "Waspada", "KLB/Siaga", "Siaga", "Waspada", "Aman"]
            })
            st.session_state.uploaded_name = "contoh_data_surveilans.csv"
            st.session_state.laporan_teks = ""
            st.rerun()

    if file_input is not None:
        try:
            if file_input.name.endswith('.csv'):
                st.session_state.df_data = pd.read_csv(file_input)
            else:
                st.session_state.df_data = pd.read_excel(file_input)
            st.session_state.uploaded_name = file_input.name
            st.session_state.laporan_teks = ""
        except Exception as err:
            st.error(f"Gagal membaca file: {err}")

    # 2. Pratinjau & Generate Laporan
    if st.session_state.df_data is not None:
        df = st.session_state.df_data
        
        with st.expander("Pratinjau Tabel Rekapitulasi Data (" + str(len(df)) + " Baris)", expanded=False):
            st.dataframe(df, use_container_width=True, height=220)
            
        col_act, col_fmt = st.columns([2, 1])
        with col_fmt:
            format_pilihan = st.selectbox(
                "Format Naskah:",
                [
                    "Laporan Resmi Kedinasan (Lengkap)",
                    "Ringkasan Eksekutif (Singkat & Padat)",
                    "Laporan Respon Cepat Lapangan (Fokus KLB)"
                ],
                index=0
            )
        with col_act:
            st.write("")
            st.write("")
            btn_proc = st.button("Buat Laporan Naratif Otomatis", type="primary", use_container_width=True)

        if btn_proc:
            with st.spinner("Memproses data dan menyusun laporan dinas resmi..."):
                try:
                    periode_ini = f"Minggu Epidemiologi ke-{datetime.now().isocalendar()[1]} ({datetime.now().strftime('%B %Y')})"
                    try:
                        table_md = df.to_markdown(index=False)
                    except Exception:
                        table_md = df.to_csv(index=False)
                    
                    if "Ringkasan Eksekutif" in format_pilihan:
                        panduan_format_khusus = f"""
STRUKTUR NASKAH RINGKASAN EKSEKUTIF (Singkat, Padat, Maksimal 3-4 Bagian):
# RINGKASAN EKSEKUTIF SURVEILANS KESEHATAN MINGGUAN
**{st.session_state.instansi.upper()}**
*Unit Pengelola: {st.session_state.bidang}*
*Periode: {periode_ini}*

---

### 1. IKHTISAR SITUASI & STATUS KEWASPADAAN
- Gambaran umum beban penyakit surveilans minggu ini secara ringkas (1-2 paragraf).

### 2. TEMUAN KUNCI & PEMETAAN WILAYAH PRIORITAS (HOTSPOT)
- Puskesmas/Kecamatan dengan kasus tertinggi (Red Zone).
- Sorotan indikator penyakit yang membutuhkan perhatian khusus.

### 3. INTERVENSI MENDESAK & REKOMENDASI TAKTIS
- 3-5 poin aksi taktis prioritas yang harus segera dieksekusi tim lapangan dan faskes.

### 4. PENGESAHAN LAPORAN
- Tempat, tanggal pelaporan ({datetime.now().strftime('%d %B %Y')}) dan slot tanda tangan Kepala Bidang/Penanggung Jawab.
"""
                    elif "Respon Cepat" in format_pilihan:
                        panduan_format_khusus = f"""
STRUKTUR NASKAH LAPORAN RESPON CEPAT LAPANGAN (Fokus Potensi KLB):
# LAPORAN RESPON CEPAT SITUASI EPIDEMIOLOGI & POTENSI KLB
**{st.session_state.instansi.upper()}**
*Unit Pengelola: {st.session_state.bidang}*
*Periode: {periode_ini}*

---

### 1. NOTIFIKASI DINI & SINYAL KEWASPADAAN KLB
- Identifikasi klaster wilayah berstatus Waspada/Siaga/KLB.

### 2. HASIL VERIFIKASI FAKTOR RISIKO LAPANGAN
- Determinan lingkungan, vektor/sanitasi, dan riwayat mobilitas kasus.

### 3. PROTOKOL PENANGGULANGAN TAKTIS (1x24 - 72 JAM)
- Penyelidikan Epidemiologi (PE), larvasidasi/fogging fokus, rujukan kasus berat, dan logistik darurat.

### 4. KOORDINASI LINTAS SEKTOR & LEMBAR PENGESAHAN
- Kebutuhan dukungan Camat/Lurah serta pengesahan penanggung jawab surveilans ({datetime.now().strftime('%d %B %Y')}).
"""
                    else:
                        panduan_format_khusus = f"""
STRUKTUR LAPORAN KEDINASAN LENGKAP:
# LAPORAN SURVEILANS EPIDEMIOLOGI MINGGUAN
**{st.session_state.instansi.upper()}**
*Unit Pengelola: {st.session_state.bidang}*
*Periode: {periode_ini}*

---

### I. LATAR BELAKANG DAN DASAR PELAKSANAAN
- Dasar hukum singkat penyelenggaraan SKDR dan tujuan surveilans periode ini.
- Cakupan wilayah dan fasilitas kesehatan yang dilaporkan.

### II. ANALISIS SITUASI DAN TREN BEBAN KASUS
- Ringkasan total agregat kasus surveilans.
- Pemetaan wilayah dengan beban kasus tertinggi (Red Zone / Hotspot) dan wilayah aman.
- Analisis perbandingan antar indikator/penyakit.

### III. IDENTIFIKASI FAKTOR RISIKO DAN KENDALA LAPANGAN
- Analisis faktor determinan (lingkungan, cuaca, mobilitas penduduk, sanitasi/PSN, pola gizi balita).
- Penilaian risiko potensi eskalasi Kejadian Luar Biasa (KLB).

### IV. REKOMENDASI TINDAKAN DAN INTERVENSI STRATEGIS
- **A. Tindakan Taktis Cepat (1-7 Hari):** Penyelidikan Epidemiologi (PE), larvasidasi/fogging fokus, rujukan faskes, droping logistik.
- **B. Tindakan Pencegahan (Jangka Menengah):** Penguatan kader jumantik, posyandu terpadu, edukasi masyarakat.
- **C. Koordinasi Lintas Sektor:** Kerja sama dengan Camat, Lurah, Dinas Lingkungan Hidup, dan aparat wilayah.

### V. KESIMPULAN DAN LEMBAR PENGESAHAN
- Ringkasan eksekutif 1 paragraf.
- Tempat dan tanggal pelaporan ({datetime.now().strftime('%d %B %Y')}).
- Kolom tanda tangan Penanggung Jawab Surveilans.
"""

                    prompt = f"""
Anda adalah Tenaga Ahli Epidemiologi dan Kepala Tim Surveilans di {st.session_state.instansi}.
Tugas Anda: Menyusun Draf Laporan Naratif Resmi Mingguan yang komprehensif, berbasis bukti, dan lugas berdasarkan data tabel berikut:

DATA SURVEILANS:
{table_md}

PARAMETER DOKUMEN:
- Instansi: {st.session_state.instansi}
- Bidang: {st.session_state.bidang}
- Periode: {periode_ini}
- Format Naskah yang Diminta: {format_pilihan}

{panduan_format_khusus}

Gunakan bahasa Indonesia formal kedinasan yang baku, terstruktur, dan tidak bertele-tele.
"""
                    generated_text = ""
                    models_to_try = [st.session_state.model_name, "gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite", "gemini-3.1-flash-lite"]
                    models_to_try = list(dict.fromkeys([m for m in models_to_try if m]))
                    last_error = None

                    for current_model in models_to_try:
                        try:
                            if HAS_NEW_GENAI:
                                client = genai.Client(api_key=st.session_state.api_key)
                                response = client.models.generate_content(
                                    model=current_model,
                                    contents=prompt
                                )
                                generated_text = response.text
                            else:
                                legacy_genai.configure(api_key=st.session_state.api_key)
                                m = legacy_genai.GenerativeModel(current_model)
                                response = m.generate_content(prompt)
                                generated_text = response.text
                            if generated_text and len(generated_text.strip()) > 0:
                                st.session_state.model_name = current_model
                                break
                        except Exception as e_model:
                            last_error = e_model
                            continue

                    if not generated_text:
                        raise last_error or Exception("Gagal menghasilkan dokumen dari AI.")

                    st.session_state.laporan_teks = generated_text
                        
                    num_cols = df.select_dtypes(include=['number']).columns.tolist()
                    metrics_str = ", ".join([f"{c}: {int(df[c].sum())}" for c in num_cols[:3]]) if num_cols else "Data Kuantitatif"
                    lines = st.session_state.laporan_teks.split("\n")
                    kesimpulan_excerpt = "Draf laporan berhasil disusun."
                    for i, l in enumerate(lines):
                        if "KESIMPULAN" in l.upper():
                            kesimpulan_excerpt = "\n".join([x for x in lines[i+1:i+6] if x.strip()])
                            break
                            
                    audit_entry = {
                        "id": f"SURV-{datetime.now().strftime('%Y%m%d-%H%M%S')}",
                        "timestamp": datetime.now().strftime("%d %B %Y, %H:%M:%S"),
                        "filename": getattr(st.session_state, "uploaded_name", "surveilans_data.csv"),
                        "total_rows": len(df),
                        "summary_metrics": metrics_str,
                        "format": format_pilihan,
                        "model_used": st.session_state.model_name,
                        "summary_highlight": f"{len(df)} Wilayah ({metrics_str})",
                        "executive_summary": kesimpulan_excerpt,
                        "full_report": st.session_state.laporan_teks
                    }
                    save_audit_log(audit_entry)
                    st.success("Laporan naratif berhasil disusun.")
                except Exception as err:
                    st.error(f"Gagal memproses laporan: {err}")

    # 3. Review & Unduh Dokumen
    if st.session_state.laporan_teks:
        st.divider()
        
        def create_docx(text_content):
            doc = Document()
            title = doc.add_heading('LAPORAN SURVEILANS KESEHATAN', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub = doc.add_paragraph(f"{st.session_state.instansi.upper()}\n{st.session_state.bidang}")
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("-" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            for line in text_content.split('\n'):
                line_str = line.strip()
                if not line_str:
                    continue
                if line_str.startswith('# '):
                    doc.add_heading(line_str.replace('# ', ''), level=1)
                elif line_str.startswith('## '):
                    doc.add_heading(line_str.replace('## ', ''), level=2)
                elif line_str.startswith('### '):
                    doc.add_heading(line_str.replace('### ', ''), level=3)
                elif line_str.startswith('- '):
                    doc.add_paragraph(line_str.replace('- ', ''), style='List Bullet')
                else:
                    doc.add_paragraph(line_str)
                    
            bio = io.BytesIO()
            doc.save(bio)
            bio.seek(0)
            return bio.getvalue()

        nama_file_base = f"Laporan_Surveilans_{datetime.now().strftime('%Y%m%d_%H%M')}"
        
        st.write("##### Unduh Dokumen Hasil Analisis:")
        col_d1, col_d2, col_d3 = st.columns(3)
        
        with col_d1:
            docx_data = create_docx(st.session_state.laporan_teks)
            st.download_button(
                label="Unduh Word (.docx)",
                data=docx_data,
                file_name=f"{nama_file_base}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                type="primary"
            )
        with col_d2:
            st.download_button(
                label="Unduh Markdown (.md)",
                data=st.session_state.laporan_teks,
                file_name=f"{nama_file_base}.md",
                mime="text/markdown",
                use_container_width=True
            )
        with col_d3:
            st.download_button(
                label="Unduh Teks Plain (.txt)",
                data=st.session_state.laporan_teks,
                file_name=f"{nama_file_base}.txt",
                mime="text/plain",
                use_container_width=True
            )

        tab_doc, tab_edit = st.tabs(["Pratinjau Dokumen Naskah Dinas", "Edit Teks Laporan"])
        
        with tab_doc:
            st.markdown(f"""
            <div class="official-paper">
                {st.session_state.laporan_teks}
            </div>
            """, unsafe_allow_html=True)
            
        with tab_edit:
            st.caption("Sesuaikan isi teks sebelum didistribusikan:")
            edit_val = st.text_area(
                "Editor:",
                value=st.session_state.laporan_teks,
                height=500,
                label_visibility="collapsed"
            )
            st.session_state.laporan_teks = edit_val

# ==============================================================================
# HALAMAN 3: PORTAL KHUSUS ADMIN (ALEX)
# ==============================================================================
elif st.session_state.active_page == "admin" and st.session_state.is_admin:
    
    col_ad1, col_ad2 = st.columns([2, 1])
    with col_ad1:
        if st.button("← Kembali ke Panduan Sistem", use_container_width=False):
            st.session_state.active_page = "tutorial"
            st.rerun()
            
    st.write("### Portal Pengawasan dan Manajemen Admin")
    st.caption("Pantau riwayat dokumen yang diproses staf dan kelola konfigurasi API Gemini.")
    
    tab_audit, tab_config = st.tabs(["Riwayat Data & Pengawasan Dokumen", "Konfigurasi API & Model AI"])
    
    with tab_audit:
        logs = load_audit_logs()
        
        col_m1, col_m2, col_m3 = st.columns(3)
        with col_m1:
            st.metric("Total Dokumen Dianalisis", len(logs))
        with col_m2:
            terakhir = logs[0]["timestamp"] if logs else "-"
            st.metric("Aktivitas Terakhir", terakhir)
        with col_m3:
            st.metric("Status Data", "Terkendali & Aman")
            
        st.write("")
        col_act1, col_act2 = st.columns([3, 1])
        with col_act1:
            st.markdown("##### Daftar Rekam Dokumen Masuk")
        with col_act2:
            if logs and st.button("Bersihkan Seluruh Log", use_container_width=True):
                clear_all_audit_logs()
                st.rerun()
                
        if not logs:
            st.info("Belum ada rekam riwayat dokumen.")
        else:
            for idx, log in enumerate(logs):
                with st.expander(f"[{log.get('timestamp')}] {log.get('filename', 'File Data')} — {log.get('summary_highlight', 'Analisis Selesai')}"):
                    st.write(f"**ID Transaksi:** `{log.get('id')}` &bull; **Model:** `{log.get('model_used')}` &bull; **Format:** `{log.get('format')}`")
                    st.write(f"**Ringkasan Kasus:** {log.get('summary_metrics', '-')}")
                    
                    st.markdown("**Kesimpulan & Rekomendasi Utama:**")
                    st.info(log.get('executive_summary', 'Tidak ada ringkasan.'))
                    
                    with st.expander("Lihat Isi Naskah Laporan Lengkap"):
                        st.text_area("Naskah:", value=log.get("full_report", ""), height=300, key=f"log_text_{idx}")
                        
                    if st.button("Hapus Rekam Ini", key=f"del_{log.get('id')}"):
                        delete_audit_log(log.get('id'))
                        st.rerun()

    with tab_config:
        st.markdown("##### Pengaturan Kunci API & Model Gemini")
        
        with st.form("form_config_admin"):
            new_api_key = st.text_input("Gemini API Key:", value=st.session_state.api_key, type="password")
            available_models = ["gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite", "gemini-3.1-flash-lite"]
            cur_idx = available_models.index(st.session_state.model_name) if st.session_state.model_name in available_models else 0
            new_model = st.selectbox(
                "Pilihan Model Gemini yang Aktif:",
                available_models,
                index=cur_idx
            )
            new_instansi = st.text_input("Nama Instansi Default:", value=st.session_state.instansi)
            new_bidang = st.text_input("Nama Bidang Default:", value=st.session_state.bidang)
            
            save_config = st.form_submit_button("Simpan Perubahan Konfigurasi", type="primary", use_container_width=True)
            
            if save_config:
                st.session_state.api_key = new_api_key
                st.session_state.model_name = new_model
                st.session_state.instansi = new_instansi
                st.session_state.bidang = new_bidang
                
                try:
                    set_key(ENV_PATH, "GEMINI_API_KEY", new_api_key)
                    set_key(ENV_PATH, "DEFAULT_MODEL", new_model)
                    set_key(ENV_PATH, "DEFAULT_INSTANSI", new_instansi)
                    set_key(ENV_PATH, "DEFAULT_BIDANG", new_bidang)
                except Exception:
                    pass
                st.success("Konfigurasi berhasil disimpan dan aktif!")
                st.rerun()
